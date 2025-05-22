import ttkbootstrap as ttk
from ttkbootstrap.constants import *
from ttkbootstrap.dialogs import Messagebox
import openpyxl
import os
import numpy as np
import pandas as pd
from tkinter import filedialog
import tkinter as tk
import matplotlib.pyplot as plt
import pathlib
import statsmodels.api as sm
from statsmodels.genmod.families import Poisson
from statsmodels.genmod.cov_struct import Exchangeable

# 定义语言字典
LANGUAGES = {
    'zh': {
        'title': "广义估计方程分析",
        'select_button': "选择文件",
        'analyze_button': "分析文件",
        'file_not_found': "文件不存在，请重新选择。",
        'analysis_success': "分析完成，结果已保存到 {}\n",
        'no_save_path': "未选择保存路径，结果未保存。",
        'analysis_error': "分析文件时出错: {}",
        'switch_language': "切换语言",
        'explanation': {
            "广义估计方程": "用于处理具有相关性的纵向数据或聚类数据，能在考虑数据相关性的情况下估计回归系数。",
        },
        'interpretation': {
            "回归系数": "表示自变量对因变量的影响程度，系数的正负表示影响方向，绝对值大小表示影响强度。",
            "标准误": "衡量回归系数估计值的抽样误差大小，标准误越小，估计越精确。",
            "p值": "若 p 值小于显著性水平（通常为 0.05），则认为该自变量对因变量有显著影响。"
        }
    },
    'en': {
        'title': "Generalized Estimating Equations Analysis",
        'select_button': "Select File",
        'analyze_button': "Analyze File",
        'file_not_found': "The file does not exist. Please select again.",
        'analysis_success': "Analysis completed. The results have been saved to {}\n",
        'no_save_path': "No save path selected. The results were not saved.",
        'analysis_error': "An error occurred while analyzing the file: {}",
        'switch_language': "Switch Language",
        'explanation': {
            "Generalized Estimating Equations": "Used to handle correlated longitudinal or clustered data, and can estimate regression coefficients while considering data correlation.",
        },
        'interpretation': {
            "Regression Coefficient": "Indicates the degree of influence of the independent variable on the dependent variable. The sign of the coefficient represents the direction of the influence, and the absolute value represents the strength of the influence.",
            "Standard Error": "Measures the sampling error of the regression coefficient estimate. A smaller standard error indicates a more precise estimate.",
            "p-value": "If the p-value is less than the significance level (usually 0.05), it is considered that the independent variable has a significant influence on the dependent variable."
        }
    }
}

# 当前语言
current_language = 'en'


def select_file():
    file_path = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx;*.xls")])
    if file_path:
        file_entry.delete(0, tk.END)
        file_entry.insert(0, file_path)
        file_entry.configure(style="TEntry")


def analyze_file():
    global current_language
    file_path = file_entry.get()
    if file_path == "请输入待分析 Excel 文件的完整路径" or file_path == "Please enter the full path of the Excel file to be analyzed":
        file_path = ""
    if not os.path.exists(file_path):
        result_label.config(text=LANGUAGES[current_language]['file_not_found'])
        return
    try:
        # 打开 Excel 文件
        df = pd.read_excel(file_path)

        # 假设第一列是聚类标识，最后一列是因变量，其余列是自变量
        cluster_id = df.iloc[:, 0]
        y = df.iloc[:, -1]
        X = df.iloc[:, 1:-1]
        X = sm.add_constant(X)

        # 进行广义估计方程分析
        fam = Poisson()
        ind = Exchangeable()
        model = sm.GEE(y, X, groups=cluster_id, cov_struct=ind, family=fam)
        result = model.fit()

        # 提取结果
        summary = result.summary()
        summary_df = pd.DataFrame(summary.tables[1].data[1:], columns=summary.tables[1].data[0])

        # 添加解释说明
        explanations = LANGUAGES[current_language]['explanation']
        interpretations = LANGUAGES[current_language]['interpretation']
        explanation_df = pd.DataFrame([explanations])
        explanation_df = explanation_df.reindex(
            columns=["广义估计方程" if current_language == 'zh' else "Generalized Estimating Equations"])
        explanation_df.insert(0, "统计量", "解释说明" if current_language == 'zh' else "Explanation")

        # 添加分析结果解读
        interpretation_df = pd.DataFrame([interpretations])
        interpretation_df = interpretation_df.reindex(columns=[
            "回归系数" if current_language == 'zh' else "Regression Coefficient",
            "标准误" if current_language == 'zh' else "Standard Error",
            "p值" if current_language == 'zh' else "p-value"
        ])
        interpretation_df.insert(0, "统计量", "结果解读" if current_language == 'zh' else "Interpretation")

        # 合并数据、解释说明和结果解读
        combined_df = pd.concat([summary_df, explanation_df, interpretation_df], ignore_index=True)

        # 让用户选择保存路径
        save_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")])
        if save_path:
            # 保存到 Excel 文件
            with pd.ExcelWriter(save_path, engine='openpyxl') as writer:
                combined_df.to_excel(writer, index=False)
                worksheet = writer.sheets['Sheet1']
                # 自动调整列宽
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = openpyxl.utils.get_column_letter(column[0].column)
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = (max_length + 2)
                    worksheet.column_dimensions[column_letter].width = adjusted_width

            # 生成结果图片（回归系数可视化）
            desktop_path = pathlib.Path.home() / 'Desktop'
            plot_path = desktop_path / 'gee_regression_coefficients.png'
            plt.figure()
            coefs = result.params[1:]
            variables = X.columns[1:]
            plt.bar(variables, coefs)
            plt.xlabel('自变量' if current_language == 'zh' else 'Independent Variables')
            plt.ylabel('回归系数' if current_language == 'zh' else 'Regression Coefficients')
            plt.title(
                '广义估计方程回归系数' if current_language == 'zh' else 'Generalized Estimating Equations Regression Coefficients')
            plt.xticks(rotation=45)
            plt.savefig(plot_path)
            plt.close()

            result_msg = LANGUAGES[current_language]['analysis_success'].format(
                save_path) + f"结果图片已保存到 {plot_path}" if current_language == 'zh' else f"The result image has been saved to {plot_path}"
            result_label.config(text=result_msg, wraplength=400)
        else:
            result_label.config(text=LANGUAGES[current_language]['no_save_path'])

    except Exception as e:
        result_label.config(text=LANGUAGES[current_language]['analysis_error'].format(str(e)))


def switch_language():
    global current_language
    current_language = 'en' if current_language == 'zh' else 'zh'
    root.title(LANGUAGES[current_language]['title'])
    select_button.config(text=LANGUAGES[current_language]['select_button'])
    analyze_button.config(text=LANGUAGES[current_language]['analyze_button'])
    language_label.config(text=LANGUAGES[current_language]['switch_language'])
    # 切换语言时更新提示信息
    file_entry.delete(0, tk.END)
    if current_language == 'zh':
        file_entry.insert(0, "请输入待分析 Excel 文件的完整路径")
        file_entry.configure(style="Gray.TEntry")
    else:
        file_entry.insert(0, "Please enter the full path of the Excel file to be analyzed")
        file_entry.configure(style="Gray.TEntry")


def on_entry_click(event):
    """当用户点击输入框时，清除提示信息"""
    if file_entry.get() == "请输入待分析 Excel 文件的完整路径" or file_entry.get() == "Please enter the full path of the Excel file to be analyzed":
        file_entry.delete(0, tk.END)
        file_entry.configure(style="TEntry")


def on_focusout(event):
    """当用户离开输入框时，如果没有输入内容，恢复提示信息"""
    if file_entry.get() == "":
        if current_language == 'zh':
            file_entry.insert(0, "请输入待分析 Excel 文件的完整路径")
            file_entry.configure(style="Gray.TEntry")
        else:
            file_entry.insert(0, "Please enter the full path of the Excel file to be analyzed")
            file_entry.configure(style="Gray.TEntry")


# 创建主窗口
root = ttk.Window(themename="flatly")
root.title(LANGUAGES[current_language]['title'])

# 获取屏幕的宽度和高度
screen_width = root.winfo_screenwidth()
screen_height = root.winfo_screenheight()

# 设置窗口的宽度和高度
window_width = 500
window_height = 300

# 计算窗口应该放置的位置
x = (screen_width - window_width) // 2
y = (screen_height - window_height) // 2

# 设置窗口的位置和大小
root.geometry(f"{window_width}x{window_height}+{x}+{y}")

# 创建自定义样式
style = ttk.Style()
style.configure("Gray.TEntry", foreground="gray")

# 创建文件选择按钮
select_button = ttk.Button(root, text=LANGUAGES[current_language]['select_button'], command=select_file,
                           bootstyle=PRIMARY)
select_button.pack(pady=10)

# 创建文件路径输入框
file_entry = ttk.Entry(root, width=50, style="Gray.TEntry")
if current_language == 'zh':
    file_entry.insert(0, "请输入待分析 Excel 文件的完整路径")
else:
    file_entry.insert(0, "Please enter the full path of the Excel file to be analyzed")
file_entry.pack(pady=5)
file_entry.bind("<FocusIn>", on_entry_click)
file_entry.bind("<FocusOut>", on_focusout)

# 创建分析按钮
analyze_button = ttk.Button(root, text=LANGUAGES[current_language]['analyze_button'], command=analyze_file,
                            bootstyle=SUCCESS)
analyze_button.pack(pady=10)

# 创建语言切换标签
language_label = ttk.Label(root, text=LANGUAGES[current_language]['switch_language'], cursor="hand2")
language_label.pack(pady=10)
language_label.bind("<Button-1>", lambda event: switch_language())

# 创建结果显示标签
result_label = ttk.Label(root, text="", justify=tk.LEFT)
result_label.pack(pady=10)

# 运行主循环
try:
    root.mainloop()
except KeyboardInterrupt:
    print("程序被手动中断。")
    root.destroy()  # 销毁主窗口
